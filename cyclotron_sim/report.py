"""Text reports and tabular exports."""

from __future__ import annotations

from pathlib import Path

import numpy as np

from .instrumentation import MeasurementResult
from .model import SimulationResult


def _value_lines(result: SimulationResult) -> list[tuple[str, str, str]]:
    config = result.config
    particle = config.particle
    return [
        ("particle", particle.label, ""),
        ("charge state", f"{particle.charge_number:+d}", "e"),
        ("mass", f"{particle.mass / 1.66053906660e-27:.4f}", "u"),
        ("magnetic field", f"{config.magnetic_field:.5f}", "T"),
        ("RF frequency", f"{config.rf_frequency / 1.0e6:.5f}", "MHz"),
        ("cyclotron frequency", f"{config.cyclotron_frequency / 1.0e6:.5f}", "MHz"),
        ("RF detuning", f"{config.detuning * 1.0e6:.2f}", "ppm"),
        ("acceleration voltage", f"{config.effective_voltage:.2f}", "V"),
        ("gap width", f"{config.gap_width * 1.0e3:.3f}", "mm"),
        ("peak gap field", f"{config.gap_field / 1.0e3:.3f}", "kV/m"),
        ("waveform", config.waveform, ""),
        ("RF phase", f"{config.rf_phase_deg:.2f}", "deg"),
        ("chamber radius", f"{config.chamber_radius * 1.0e3:.2f}", "mm"),
        ("start position", f"({config.x0 * 1.0e3:.2f}, {config.y0 * 1.0e3:.2f})", "mm"),
        ("start velocity", f"({config.vx0:.3e}, {config.vy0:.3e})", "m/s"),
        ("Faraday cup", "inserted" if config.cup_enabled else "not inserted", ""),
        ("cup radius", f"{config.cup_radius_mm:.2f}", "mm"),
    ]


def parameter_rows(result: SimulationResult) -> list[tuple[str, str, str]]:
    return _value_lines(result)


def result_rows(
    result: SimulationResult, measurement: MeasurementResult
) -> list[tuple[str, str, str]]:
    config = result.config
    measured_frequency = measurement.pickup_frequency_hz
    frequency_resolved = bool(np.isfinite(measured_frequency))
    frequency_error = (
        (measured_frequency - config.cyclotron_frequency)
        / config.cyclotron_frequency
        * 100.0
        if frequency_resolved and config.cyclotron_frequency
        else np.nan
    )
    energy_error = (
        (measurement.measured_final_energy_kev - result.final_energy_kev)
        / result.final_energy_kev
        * 100.0
        if result.final_energy_kev
        else np.nan
    )
    beta = result.final_speed / 2.99792458e8
    frequency_text = (
        f"{measured_frequency / 1.0e6:.5f}" if frequency_resolved else "not resolved"
    )
    uncertainty_text = (
        f"{measurement.pickup_frequency_uncertainty_hz / 1.0e3:.3f}"
        if np.isfinite(measurement.pickup_frequency_uncertainty_hz)
        else "n/a"
    )
    frequency_error_text = (
        f"{frequency_error:+.3f}" if frequency_resolved else "n/a"
    )
    fft_text = (
        f"{measurement.fft_frequency_hz / 1.0e6:.5f}"
        if np.isfinite(measurement.fft_frequency_hz)
        else "not resolved"
    )
    gain_text = (
        f"{result.energy_gain_per_turn_kev:.4f}"
        if result.n_turns >= 2
        else "n/a"
    )
    b_from_pickup = (
        2.0
        * np.pi
        * config.particle.mass
        * measured_frequency
        / abs(config.particle.charge)
        if frequency_resolved
        else np.nan
    )
    b_error = (
        (b_from_pickup - config.magnetic_field)
        / config.magnetic_field
        * 100.0
        if frequency_resolved
        else np.nan
    )
    b_text = f"{b_from_pickup:.5f}" if frequency_resolved else "n/a"
    b_error_text = f"{b_error:+.3f}" if frequency_resolved else "n/a"
    return [
        ("exit condition", result.exit_reason, ""),
        ("simulated time", f"{result.duration * 1.0e6:.4f}", "us"),
        ("complete turns", f"{result.n_turns:d}", ""),
        ("final speed", f"{result.final_speed / 1.0e3:.3f}", "km/s"),
        ("final speed / c", f"{beta:.5f}", ""),
        ("final energy", f"{result.final_energy_kev:.4f}", "keV"),
        ("final radius", f"{result.final_radius_mm:.4f}", "mm"),
        ("energy gain per turn", gain_text, "keV"),
        ("mean crossing phase", f"{np.rad2deg(result.mean_crossing_phase):.2f}", "deg"),
        ("path length", f"{result.total_path_length:.5f}", "m"),
        ("gap crossings", f"{result.gap_crossings_t.size:d}", ""),
        ("pickup frequency", frequency_text, "MHz"),
        ("pickup frequency uncertainty", uncertainty_text, "kHz"),
        ("pickup frequency error", frequency_error_text, "%"),
        ("B from pickup frequency", b_text, "T"),
        ("B measurement error", b_error_text, "%"),
        ("FFT frequency", fft_text, "MHz"),
        (
            "energy from B-rho",
            f"{measurement.measured_final_energy_kev:.4f}",
            "keV",
        ),
        (
            "energy measurement uncertainty",
            f"{measurement.measured_energy_uncertainty_kev:.4f}",
            "keV",
        ),
        ("energy error", f"{energy_error:+.3f}", "%"),
        (
            "Faraday cup charge",
            f"{measurement.cup_charge_c * 1.0e12:.4f}",
            "pC",
        ),
        (
            "Faraday cup peak current",
            f"{measurement.cup_peak_current_a * 1.0e6:.4f}",
            "uA",
        ),
        ("Faraday cup hits", f"{result.cup_hit_times.size:d}", ""),
    ]


def turn_rows(result: SimulationResult) -> list[tuple[float, float, float, float]]:
    rows = []
    for index, (time, radius, energy) in enumerate(
        zip(result.turn_times, result.turn_radii, result.turn_energies_j), start=1
    ):
        rows.append(
            (
                index,
                time * 1.0e6,
                radius * 1.0e3,
                energy / 1.602176634e-16,
            )
        )
    return rows


def crossing_rows(result: SimulationResult) -> list[tuple[int, float, float, float]]:
    rows = []
    for index, (time, phase, energy) in enumerate(
        zip(
            result.gap_crossings_t,
            result.gap_crossings_phase,
            result.gap_crossings_energy_j,
        ),
        start=1,
    ):
        rows.append(
            (
                index,
                time * 1.0e6,
                np.rad2deg(phase),
                energy / 1.602176634e-16,
            )
        )
    return rows


def format_report(
    result: SimulationResult, measurement: MeasurementResult
) -> str:
    width = 78
    lines = [
        "=" * width,
        "CYCLOTRON SIMULATION AND INSTRUMENTATION REPORT",
        "=" * width,
        "",
        "OPERATING POINT",
        "-" * width,
    ]
    for name, value, unit in _value_lines(result):
        lines.append(f"{name:<28}{value:>30}  {unit}".rstrip())

    lines.extend(["", "SIMULATED BEAM", "-" * width])
    for name, value, unit in result_rows(result, measurement):
        lines.append(f"{name:<34}{value:>24}  {unit}".rstrip())

    if result.turn_times.size:
        lines.extend(["", "TURN-BY-TURN DATA", "-" * width])
        lines.append(f"{'turn':>5}{'time (us)':>14}{'radius (mm)':>15}{'energy (keV)':>16}")
        for turn, time, radius, energy in turn_rows(result):
            lines.append(
                f"{turn:>5}{time:>14.4f}{radius:>15.4f}{energy:>16.4f}"
            )

    lines.extend(
        [
            "",
            "CROSSING PHASES",
            "-" * width,
            f"{'crossing':>9}{'time (us)':>14}{'RF phase (deg)':>17}{'dE (keV)':>14}",
        ]
    )
    for crossing, time, phase, energy in crossing_rows(result)[:24]:
        lines.append(f"{crossing:>9}{time:>14.4f}{phase:>17.2f}{energy:>14.4f}")
    if len(result.gap_crossings_t) > 24:
        lines.append(f"... {len(result.gap_crossings_t) - 24} more crossings")

    lines.extend(
        [
            "",
            "MEASUREMENT NOTES",
            "-" * width,
            f"sampling rate             {measurement.instrument.sample_rate_hz / 1.0e9:.2f} GS/s",
            f"pickup noise              {measurement.instrument.pickup_noise_rms_v * 1.0e3:.3f} mV rms",
            f"timing jitter             {measurement.instrument.timing_jitter_s * 1.0e12:.2f} ps rms",
            f"beam particles per bunch  {measurement.instrument.beam_particles:.3e}",
            f"radius resolution         {measurement.instrument.radius_resolution_mm:.3f} mm rms",
            f"field relative error      {measurement.instrument.field_relative_error * 100.0:.3f} %",
            "",
        ]
    )
    return "\n".join(lines)


def write_excel(
    result: SimulationResult,
    measurement: MeasurementResult,
    path: str | Path,
) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Cyclotron simulation report"])
    summary.append([])
    summary.append(["Parameter", "Value", "Unit"])
    for row in _value_lines(result):
        summary.append(list(row))
    summary.append([])
    summary.append(["Result", "Value", "Unit"])
    for row in result_rows(result, measurement):
        summary.append(list(row))

    turns = workbook.create_sheet("Turns")
    turns.append(["turn", "time_us", "radius_mm", "energy_keV"])
    for row in turn_rows(result):
        turns.append(list(row))

    crossings = workbook.create_sheet("Crossings")
    crossings.append(["crossing", "time_us", "rf_phase_deg", "energy_gain_keV"])
    for row in crossing_rows(result):
        crossings.append(list(row))

    signals = workbook.create_sheet("Signals")
    signals.append(["time_us", "rf_voltage_V", "pickup_mV", "cup_uA"])
    for index in range(measurement.t.size):
        signals.append(
            [
                measurement.t[index] * 1.0e6,
                float(measurement.rf_voltage_v[index]),
                float(measurement.pickup_voltage_v[index] * 1.0e3),
                float(measurement.cup_current_a[index] * 1.0e6),
            ]
        )

    spectrum = workbook.create_sheet("Spectrum")
    spectrum.append(["frequency_MHz", "amplitude_mV"])
    for frequency, amplitude in zip(
        measurement.fft_frequencies_hz, measurement.fft_spectrum
    ):
        spectrum.append([frequency * 1.0e-6, float(amplitude * 1.0e3)])

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def write_csv_files(
    result: SimulationResult,
    measurement: MeasurementResult,
    output_dir: str | Path,
) -> list[Path]:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)

    trajectory = output / "trajectory.csv"
    np.savetxt(
        trajectory,
        np.column_stack(
            (
                result.t,
                result.x,
                result.y,
                result.vx,
                result.vy,
                result.speed,
                result.radius * 1.0e3,
                result.energy_j / 1.602176634e-16,
            )
        ),
        delimiter=",",
        header="time_s,x_m,y_m,vx_m_s,vy_m_s,speed_m_s,radius_mm,energy_keV",
        comments="",
    )

    signals = output / "signals.csv"
    np.savetxt(
        signals,
        np.column_stack(
            (
                measurement.t,
                measurement.rf_voltage_v,
                measurement.pickup_voltage_v * 1.0e3,
                measurement.cup_current_a * 1.0e6,
            )
        ),
        delimiter=",",
        header="time_s,rf_voltage_V,pickup_mV,cup_uA",
        comments="",
    )
    return [trajectory, signals]
